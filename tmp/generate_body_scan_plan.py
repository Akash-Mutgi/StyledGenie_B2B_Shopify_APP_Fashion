from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

DEPS = Path(__file__).resolve().parent / "docx_deps"
sys.path.insert(0, str(DEPS))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "StyledGenie_Full_Body_Onboarding_Scan_Product_Matching_Plan.docx"

# compact_reference_guide preset, resolved to exact values.
FONT = "Calibri"
BODY_SIZE = 11
INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "59636E"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
GREEN_FILL = "EAF4EA"
GOLD_FILL = "FFF4D6"
BORDER = "B9C4CF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", CELL_MARGIN_TOP_BOTTOM),
        ("bottom", CELL_MARGIN_TOP_BOTTOM),
        ("start", CELL_MARGIN_START_END),
        ("end", CELL_MARGIN_START_END),
    ):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    assert sum(widths_dxa) == TABLE_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=BODY_SIZE, color="000000", bold=None, italic=None, name=FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", 1)
    code.font.name = "Courier New"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.22)
    code.paragraph_format.right_indent = Inches(0.22)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0
    p_pr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_FILL)
    p_pr.append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("StyledGenie  |  Product & AI Implementation Guide")
    set_run_font(hr, size=9, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_para(doc, text: str, *, bold_lead: str | None = None, italic=False, color="000000"):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, italic=italic, color=color)
    return p


def add_bullets(doc, items: list[str]) -> None:
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(text)
        set_run_font(run)


def add_numbers(doc, items: list[str]) -> None:
    for text in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(text)
        set_run_font(run)


def add_callout(doc, label: str, text: str, fill: str = GOLD_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_table_borders(table, color=fill, size=4)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, bold=True, color=INK)
    r2 = p.add_run(text)
    set_run_font(r2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(header) < 15 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        set_run_font(r, size=9.5, bold=True, color=INK)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    r = p.add_run(text)
    set_run_font(r, size=8.5, name="Courier New")


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_page(doc)

    # Opening block: customer_pack pattern with a project-specific title override.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("STYLEDGENIE IMPLEMENTATION GUIDE")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Full-Body Onboarding Scan")
    set_run_font(r, size=30, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("From a consented photo to the best validated Shopify product match")
    set_run_font(r, size=15, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("Project: StyledGenie B2B Shopify Fashion App  |  Prepared: 27 August 2026  |  Status: Build plan")
    set_run_font(r, size=9.5, color=MUTED, italic=True)

    add_callout(
        doc,
        "Outcome",
        "Add an image-first onboarding scan that validates full-body visibility, produces an editable fit profile, and ranks only purchasable catalog variants that are compatible with the shopper's confirmed profile.",
        fill=GREEN_FILL,
    )
    add_para(
        doc,
        "Important accuracy boundary: one uncalibrated 2D photo cannot guarantee exact physical fit. In StyledGenie, 'exact product' should mean the highest-confidence, in-stock catalog match after pose quality, confirmed shopper inputs, and garment size data are all validated.",
        bold_lead="Important accuracy boundary:",
    )

    add_heading(doc, "1. Executive recommendation", 1)
    add_para(doc, "Keep the current Google Vision -> OpenAI scan pipeline and add a client-side pose-quality gate before upload. Then extend the catalog and recommendation layer with variant-level garment measurements and fit scoring.")
    add_bullets(doc, [
        "Required addition: MediaPipe Pose Landmarker in the storefront widget for live full-body framing and landmark confidence.",
        "Required catalog work: store every sellable variant's size, garment measurements, fit type, stretch, and availability.",
        "Required UX: show estimates as editable, ask only for missing critical data, and include one short 'Why this works' explanation.",
        "Required safety: use pose geometry for fit guidance only; do not perform identity recognition or infer sensitive traits.",
        "Recommended storage model: process the scan in memory and persist the confirmed fit profile, not the raw body photo, unless the shopper separately opts in.",
    ])

    add_heading(doc, "2. What the project already has", 1)
    add_para(doc, "The repository already contains the core scan path. The implementation should extend it rather than replace it.")
    add_table(doc, ["Layer", "Current implementation", "Assessment"], [
        ["Onboarding UI", "apps/storefront-widget/app.js captures or uploads a photo and posts it to /api/onboarding/analyze-scan.", "Keep; add live pose framing and quality states."],
        ["API route", "backend/app/routers/chat.py exposes POST /api/onboarding/analyze-scan.", "Keep route; extend request/response schemas."],
        ["Orchestration", "ConversationService runs Google Vision first, then the OpenAI onboarding analyzer.", "Correct image-first order; preserve it."],
        ["Visual evidence", "VisionService requests labels, objects, image properties, web entities, and OCR.", "Good for fashion/color cues; not a body-pose validator."],
        ["AI estimate", "OpenAI returns full_body_visible, skin-tone band, body-shape class, confidence, and a quality note.", "Keep editable estimates; add structured pose evidence."],
        ["Catalog", "Supabase products store one Shopify variant ID plus product/card and inventory fields.", "Insufficient for exact size/fit matching."],
    ], [1750, 4850, 2760])

    add_heading(doc, "3. Target user flow", 1)
    add_numbers(doc, [
        "Consent: explain what is analyzed, what is saved, and provide a manual-entry alternative.",
        "Capture: open the camera over HTTPS and show a full-body guide with head and feet inside frame.",
        "Validate on device: MediaPipe checks pose visibility, orientation, framing, blur/lighting signals, and landmark confidence.",
        "Analyze image-first: send an accepted frame to the existing endpoint; Google Vision runs before OpenAI reasoning.",
        "Review: show skin-tone band and body-shape/fit cues as editable estimates, with a confidence label and one retake action if needed.",
        "Confirm sizing: collect only missing critical inputs such as height, usual top/bottom size, fit preference, and optional measurements.",
        "Match catalog: hard-filter by segment, category, availability, budget, and compatible size; then rank remaining variants.",
        "Decide: show one best product when decision mode is selected, otherwise show a small set of validated options.",
    ])
    add_callout(doc, "UX rule", "Do not turn onboarding into a questionnaire. Infer first, keep estimates editable, and ask only for data required to improve fit confidence.")

    add_heading(doc, "4. Required APIs and components", 1)
    add_table(doc, ["API / component", "Role in StyledGenie", "Need"], [
        ["Browser MediaDevices (getUserMedia)", "Camera capture in onboarding. Requires HTTPS and user permission.", "Required"],
        ["MediaPipe Pose Landmarker for Web", "33 body landmarks, optional segmentation mask, real-time pose/framing quality. Runs client-side; it is not a Google Cloud API toggle.", "Required"],
        ["Google Cloud Vision API", "Existing labels, object localization, dominant colors, web cues, and OCR before AI reasoning.", "Already integrated"],
        ["OpenAI Responses API with image input", "Existing structured reasoning for editable fashion attributes and recommendation explanations.", "Already integrated"],
        ["Shopify Admin GraphQL API", "Catalog sync for products, variants, size options, inventory, metafields, images, and prices.", "Extend current sync"],
        ["Supabase/Postgres", "Persist normalized catalog fit data and the shopper's confirmed fit profile.", "Extend schema"],
        ["Cloud Storage", "Raw scan retention only if business/legal need and explicit opt-in exist.", "Optional; avoid by default"],
        ["Vertex AI/custom model", "Only for a future trained measurement or fit model with an evaluated dataset.", "Not needed for MVP"],
    ], [2550, 5200, 1610])

    add_heading(doc, "5. Capture and scan-quality contract", 1)
    add_para(doc, "The photo should be accepted only when the capture gate passes. This prevents the AI from guessing from partial or unsuitable images.")
    add_bullets(doc, [
        "Person count equals one; shoulders, hips, knees, ankles, and feet are visible.",
        "Front-facing neutral pose; arms slightly separated from the torso and feet approximately hip-width apart.",
        "Camera is near waist/chest height, not tilted, and the full body occupies a useful portion of the frame.",
        "Lighting is even enough to inspect the outline and visible skin without harsh backlighting.",
        "Clothing is close enough to the body to estimate silhouette; loose coats or heavy layers trigger a caution or manual path.",
        "Required landmarks meet a configurable presence/visibility threshold for several consecutive frames.",
        "No automatic capture until consent has been accepted; provide retake and manual-entry controls at all times.",
    ])
    add_code(doc, "Suggested client gate (initial values; calibrate with real devices)\nrequiredLandmarkVisibility >= 0.75\nrequiredLandmarksPassing >= 12 of 14\nfullBodyInsideGuide = true\npersonCount = 1\nstableFrames >= 8\nresult: pass | caution | retake")

    add_heading(doc, "6. Scan response and confirmed profile", 1)
    add_para(doc, "Extend the current OnboardingScanAnalysisResponse instead of creating a second scan API. Keep image-derived values separate from user-confirmed fields so downstream matching knows which data is trustworthy.")
    add_code(doc, "{\n  \"full_body_visible\": true,\n  \"pose_quality\": {\"status\": \"pass\", \"score\": 0.89},\n  \"skin_tone_index\": 3,\n  \"body_shape\": \"rectangle\",\n  \"proportion_cues\": {\"shoulder_to_hip_ratio\": 1.04},\n  \"confidence\": \"medium\",\n  \"quality_note\": \"Review the estimate before saving.\",\n  \"requires_confirmation\": true,\n  \"vision_source\": \"google-vision-client\"\n}")
    add_para(doc, "Persist the confirmed fit profile with provenance:")
    add_bullets(doc, [
        "segment preference; height band; usual top, bottom, and shoe sizes; size systems; and preferred fit.",
        "optional self-entered chest/bust, waist, hip, inseam, and shoulder measurements.",
        "body-shape and proportion cues only after review; confidence and source for every inferred field.",
        "photo_retained = false by default; consent version, timestamp, and deletion status.",
        "Do not store pose landmarks as an identity template or use them to recognize a person.",
    ])

    add_heading(doc, "7. Catalog data required for accurate product matching", 1)
    add_para(doc, "The current products table does not contain a size chart or variant-level garment measurements. This is the primary data gap. Add a product_variants table plus structured fit attributes or equivalent Shopify metafields.")
    add_table(doc, ["Field group", "Minimum fields", "Why it matters"], [
        ["Variant identity", "product_id, shopify_variant_id, SKU, option values, size label, size system", "Select the actual purchasable size rather than a generic product."],
        ["Commerce", "available_for_sale, inventory_quantity/policy, price, currency, product URL", "Hard-filter unavailable or non-cart-ready variants."],
        ["Garment measures", "chest/bust, waist, hip, inseam, rise, shoulder, sleeve, length; unit", "Compare shopper measurements with the garment, including ease."],
        ["Fit behavior", "fit type, cut, stretch level/percentage, fabric, intended ease", "Two garments with the same label can fit differently."],
        ["Styling", "segment, category, silhouette, color family, occasion, weather, style tags", "Preserve StyledGenie's fashion reasoning and mode separation."],
        ["Quality", "measurement source, completeness, last_synced_at", "Lower confidence when merchant fit data is incomplete or stale."],
    ], [1800, 4500, 3060])

    add_heading(doc, "8. Product matching logic", 1)
    add_heading(doc, "8.1 Hard filters", 2)
    add_bullets(doc, [
        "Active styling mode only; support mode must never call the styling matcher.",
        "Selected/inferred shopping segment, requested category or Complete My Look gap, and budget.",
        "available_for_sale = true and a valid cart-ready Shopify variant ID.",
        "A compatible size/measurement range exists; otherwise mark fit confidence low or exclude.",
        "Merchant/store scope and shipping eligibility when available.",
    ])
    add_heading(doc, "8.2 Ranking after filtering", 2)
    add_para(doc, "Recommended initial score (100 points):")
    add_bullets(doc, [
        "30 points - size and garment-measurement compatibility, including desired ease.",
        "20 points - silhouette compatibility with confirmed fit goals and proportion cues.",
        "15 points - category/gap relevance to the active styling feature.",
        "15 points - occasion and weather fit.",
        "10 points - color harmony with preferences or an uploaded anchor item.",
        "10 points - style identity and experimentation preference.",
    ])
    add_code(doc, "fit_score = 0.30*size_compatibility + 0.20*silhouette +\n            0.15*category_gap + 0.15*occasion_weather +\n            0.10*color_harmony + 0.10*style_identity")
    add_para(doc, "Return fit_confidence (high/medium/low) separately from style_score. A stylish item with missing size data must not be presented as an exact fit.")

    add_heading(doc, "9. Project-specific code change map", 1)
    add_table(doc, ["File / area", "Minimal change"], [
        ["apps/storefront-widget/app.js", "Add camera stream, MediaPipe loader/helper, landmark quality gate, capture overlay state, and pose_quality payload; keep the existing finishOnboardingScanCapture flow."],
        ["apps/storefront-widget/styles.css", "Add responsive scan guide, pass/caution/retake states, privacy text, and accessible status messaging."],
        ["backend/app/models/schemas.py", "Extend the scan request/response and ShopperProfile with pose quality, proportion cues, source, confirmation, and fit fields."],
        ["backend/app/routers/chat.py", "Keep POST /api/onboarding/analyze-scan; accept the extended payload and retain response validation."],
        ["backend/app/services/conversation_service.py", "Pass pose evidence into the existing Vision-first -> OpenAI analyzer path."],
        ["backend/app/services/openai_service.py", "Use pose evidence as geometry support; preserve current sensitive-attribute prohibitions and editable outputs."],
        ["backend/app/services/shopper_profile_service.py", "Merge only confirmed scan fields; preserve provenance and manual overrides."],
        ["backend/app/services/shopify_service.py", "Sync all variants and normalized fit/measurement metafields, not only the first variant."],
        ["backend/app/services/supabase_service.py", "Read/write variant fit data and confirmed profile fields."],
        ["backend/app/services/recommendation_service.py", "Add availability/size hard filters, fit score, confidence, and segment leakage tests."],
        ["supabase migrations", "Add product_variants, product_fit_attributes, and shopper_fit_profiles (or equivalent normalized tables)."],
    ], [2900, 6460])

    add_heading(doc, "10. Privacy, security, and responsible-use rules", 1)
    add_bullets(doc, [
        "Obtain explicit scan consent before camera access; clearly explain the purpose and manual alternative.",
        "Run pose detection on device where practical and upload only the accepted frame plus non-identifying quality metrics.",
        "Do not identify the shopper or infer race, ethnicity, age, health, disability, attractiveness, or gender identity.",
        "Do not expose OpenAI, Google, Shopify, Supabase, or AWS secrets in the widget. Keep secrets in Elastic Beanstalk environment properties or a managed secret store.",
        "Use HTTPS end-to-end. Browser camera APIs will fail or be restricted on insecure production origins.",
        "Default to no raw-photo retention. If retention is introduced, require separate opt-in, short expiry, encryption, access logging, and a deletion endpoint.",
        "Never log image base64, camera frames, secrets, or detailed body measurements in application logs.",
    ])
    add_callout(doc, "Security action", "The API key visible in the current IDE context should be revoked and replaced before any deployment. Store the replacement only in protected environment configuration, never in screenshots, chat, source control, or client-side JavaScript.", fill=GOLD_FILL)

    add_heading(doc, "11. AWS and deployment notes", 1)
    add_bullets(doc, [
        "Bundle or reliably host the MediaPipe WebAssembly/model assets and allow only the required sources in Content-Security-Policy.",
        "Terminate TLS with the existing AWS load balancer/CloudFront path and redirect HTTP to HTTPS.",
        "Set Google credentials, OpenAI key, Shopify credentials, and Supabase secrets as Elastic Beanstalk environment properties or AWS Secrets Manager references.",
        "Do not add a Google Cloud API merely for MediaPipe; the Web Pose Landmarker runs in the browser.",
        "Set request body limits, timeouts, and image compression so mobile captures do not exhaust the application worker.",
        "Add health checks that report configured/not-configured status without returning secret values.",
    ])

    add_heading(doc, "12. Delivery phases", 1)
    add_table(doc, ["Phase", "Deliverable", "Exit gate"], [
        ["0 - Baseline", "Instrument current scan success, retake rate, and response time.", "Known baseline and test devices."],
        ["1 - Capture quality", "Camera flow, consent, MediaPipe pose gate, overlay, retake/manual path.", "Full-body acceptance works across target mobile browsers."],
        ["2 - Confirmed profile", "Extended scan schema, editable review, provenance, no-photo-retention default.", "No low-confidence estimate is silently saved."],
        ["3 - Catalog fit data", "Variant sync, size/measurement metafields, normalized Supabase tables.", "Most sellable apparel variants have usable size data."],
        ["4 - Fit ranking", "Hard filters, fit score, confidence, Why this works, decision mode.", "Unavailable/wrong-segment/wrong-size leakage tests pass."],
        ["5 - Evaluation", "Offline labeled scenarios, mobile usability, privacy review, analytics.", "Approved thresholds and rollback plan."],
    ], [1320, 4880, 3160])

    add_heading(doc, "13. Acceptance criteria", 1)
    add_bullets(doc, [
        "A partial-body or multi-person image is rejected before AI estimation, with one short retake instruction.",
        "A valid scan reaches Google Vision first, then OpenAI, and all estimated fields remain editable.",
        "Male/female/segment profile selection is preserved through scan, recommendation, and smart swap; no cross-segment leakage appears.",
        "Only in-stock, cart-ready Shopify variants with compatible size evidence can receive high fit confidence.",
        "Decision mode returns one best product; options mode returns a small ranked set.",
        "Every recommendation includes a short Why this works covering color, silhouette, occasion, and style direction where evidence exists.",
        "Smart swap replaces only the selected category and preserves all other outfit context.",
        "Support intent disables scan/styling recommendation UI completely.",
        "Raw image/base64 and secrets never appear in logs, analytics, or client bundles.",
        "The manual onboarding path remains fully usable when camera permission is denied or model loading fails.",
    ])

    add_heading(doc, "14. Test scenarios", 1)
    add_table(doc, ["Scenario", "Expected result"], [
        ["Head or feet outside frame", "Retake; no body-shape estimate saved."],
        ["Loose coat obscures outline", "Caution and manual measurement option; confidence reduced."],
        ["Valid male profile scan", "Only compatible menswear/selected-segment variants appear."],
        ["Valid female profile scan", "Only compatible womenswear/selected-segment variants appear."],
        ["No compatible size in stock", "Do not claim exact match; show the nearest validated alternative or a clear no-match action."],
        ["Camera denied/offline model", "Manual onboarding continues; no dead end."],
        ["Support request during onboarding", "Switch to support mode and remove styling controls."],
        ["Swap one jacket", "Only jacket changes; bottoms, shoes, accessories, profile, and context remain fixed."],
    ], [3150, 6210])

    add_heading(doc, "15. Implementation checklist", 1)
    add_bullets(doc, [
        "Define the Shopify metafield namespace and merchant size-data requirements.",
        "Add MediaPipe pose validation behind a feature flag.",
        "Extend schemas and persistence with field provenance and confirmation status.",
        "Change Shopify sync from first-variant storage to all-variant storage.",
        "Implement hard availability/segment/size filters before AI ranking.",
        "Add fit confidence, Why this works, and no-match behavior to product cards.",
        "Create automated tests for segment leakage, unavailable variants, scan failure, and smart swap context preservation.",
        "Complete privacy review, rotate exposed credentials, configure AWS secrets, and verify HTTPS camera behavior.",
    ])

    add_heading(doc, "16. Reference links", 1)
    add_para(doc, "Google Cloud Vision feature list: https://docs.cloud.google.com/vision/docs/features-list")
    add_para(doc, "MediaPipe Pose Landmarker: https://developers.google.com/edge/mediapipe/solutions/vision/pose_landmarker")
    add_para(doc, "OpenAI Responses API: https://developers.openai.com/api/reference/cli/resources/responses/methods/create")
    add_para(doc, "Repository evidence used: apps/storefront-widget/app.js; backend/app/routers/chat.py; backend/app/services/conversation_service.py; backend/app/services/vision_service.py; backend/app/services/openai_service.py; backend/app/models/schemas.py; backend/app/services/shopify_service.py; backend/app/services/recommendation_service.py; supabase/schema.sql.", italic=True, color=MUTED)

    core = doc.core_properties
    core.title = "StyledGenie Full-Body Onboarding Scan and Product Matching Plan"
    core.subject = "Project-specific implementation plan for image-first onboarding and fit-aware Shopify product matching"
    core.author = "StyledGenie Project Team"
    core.keywords = "StyledGenie, onboarding, body scan, MediaPipe, Google Vision, OpenAI, Shopify, product matching"
    core.comments = "Generated as a project implementation guide. No secrets or shopper images are embedded."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
