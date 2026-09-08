from __future__ import annotations

import sys
import zipfile
from pathlib import Path

DEPS = Path(__file__).resolve().parent / "docx_deps"
sys.path.insert(0, str(DEPS))

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(__file__).resolve().parents[1] / "output" / "StyledGenie_Full_Body_Onboarding_Scan_Product_Matching_Plan.docx"


def main() -> None:
    failures: list[str] = []
    with zipfile.ZipFile(DOCX) as package:
        bad = package.testzip()
        if bad:
            failures.append(f"Corrupt package member: {bad}")
        names = set(package.namelist())
        for required in {"word/document.xml", "word/styles.xml", "word/numbering.xml", "docProps/core.xml"}:
            if required not in names:
                failures.append(f"Missing package member: {required}")
        package_text = b"\n".join(package.read(name) for name in names if name.endswith((".xml", ".rels")))
        if b"sk-proj-" in package_text:
            failures.append("Secret-looking OpenAI project key embedded in DOCX")
        if any(name.startswith("word/media/") for name in names) or b"data:image/" in package_text.lower():
            failures.append("Embedded image payload found in DOCX")

    doc = Document(DOCX)
    section = doc.sections[0]
    expected = {
        "page_width": 8.5,
        "page_height": 11.0,
        "top_margin": 1.0,
        "right_margin": 1.0,
        "bottom_margin": 1.0,
        "left_margin": 1.0,
        "header_distance": 0.492,
        "footer_distance": 0.492,
    }
    for field, value in expected.items():
        actual = getattr(section, field).inches
        if abs(actual - value) > 0.01:
            failures.append(f"{field}: expected {value}, got {actual}")

    heading_count = sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading"))
    list_count = sum(1 for p in doc.paragraphs if p.style.name in {"List Bullet", "List Number"})
    if heading_count < 16:
        failures.append(f"Too few headings: {heading_count}")
    if list_count < 45:
        failures.append(f"Too few list items: {list_count}")
    if len(doc.tables) < 8:
        failures.append(f"Too few tables: {len(doc.tables)}")

    for table_index, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        width = tbl_w.get(qn("w:w")) if tbl_w is not None else None
        indent = tbl_ind.get(qn("w:w")) if tbl_ind is not None else None
        if width != "9360":
            failures.append(f"Table {table_index} width is {width}, expected 9360")
        if indent != "120":
            failures.append(f"Table {table_index} indent is {indent}, expected 120")
        grid_widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        if sum(grid_widths) != 9360:
            failures.append(f"Table {table_index} grid sums to {sum(grid_widths)}")
        for row_index, row in enumerate(table.rows, start=1):
            if row_index == 1 and (len(table.rows) > 1 or len(table.columns) > 1):
                tr_pr = row._tr.trPr
                if tr_pr is None or tr_pr.find(qn("w:tblHeader")) is None:
                    failures.append(f"Table {table_index} missing repeated/header row marker")
            for cell_index, cell in enumerate(row.cells):
                expected_width = grid_widths[min(cell_index, len(grid_widths) - 1)]
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                actual_width = int(tc_w.get(qn("w:w"))) if tc_w is not None else -1
                if actual_width != expected_width:
                    failures.append(
                        f"Table {table_index} row {row_index} cell {cell_index + 1} width {actual_width}, expected {expected_width}"
                    )
                if not "".join(p.text for p in cell.paragraphs).strip():
                    failures.append(f"Table {table_index} row {row_index} cell {cell_index + 1} is empty")

    required_sections = {
        "1. Executive recommendation",
        "4. Required APIs and components",
        "7. Catalog data required for accurate product matching",
        "9. Project-specific code change map",
        "10. Privacy, security, and responsible-use rules",
        "13. Acceptance criteria",
    }
    paragraph_texts = {p.text.strip() for p in doc.paragraphs}
    missing = sorted(required_sections - paragraph_texts)
    if missing:
        failures.append(f"Missing required sections: {missing}")

    print(f"File: {DOCX}")
    print(f"Size: {DOCX.stat().st_size} bytes")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Headings: {heading_count}")
    print(f"List items: {list_count}")
    print(f"Tables: {len(doc.tables)}")
    if failures:
        print("FAIL")
        for item in failures:
            print(f"- {item}")
        raise SystemExit(1)
    print("PASS: package, content, geometry, table headers, and secret scan checks")


if __name__ == "__main__":
    main()
